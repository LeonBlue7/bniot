"""
Excel导出服务模块
Phase 3.3 数据管理

提供数据导出为Excel格式的核心功能
"""
import io
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


class ExcelExportService:
    """Excel导出服务"""

    def __init__(self):
        self.workbook = None

    def export_data(
        self,
        data: list[dict[str, Any]],
        headers: list[str],
        sheet_name: str = "Sheet1",
        style_header: bool = True,
        number_format: str | None = None
    ) -> bytes:
        """
        导出数据为Excel格式

        Args:
            data: 数据列表（字典列表）
            headers: 表头列表
            sheet_name: 工作表名称
            style_header: 是否应用表头样式
            number_format: 数值格式（可选）

        Returns:
            Excel文件字节流
        """
        # 创建工作簿
        self.workbook = Workbook()
        sheet = self.workbook.active
        sheet.title = sheet_name

        # 写入表头
        for col_idx, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_idx, value=header)

            if style_header:
                # 应用表头样式
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                cell.border = Border(
                    bottom=Side(style="thin"),
                    right=Side(style="thin")
                )

        # 写入数据
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, header in enumerate(headers, 1):
                # 从字典中获取值（如果不存在则为None）
                value = row_data.get(header, None) if isinstance(row_data, dict) else None

                # 处理特殊类型
                if isinstance(value, datetime):
                    value = value.strftime("%Y-%m-%d %H:%M:%S")
                elif isinstance(value, bool):
                    value = "是" if value else "否"
                elif value is None:
                    value = ""

                cell = sheet.cell(row=row_idx, column=col_idx, value=value)

                # 应用数值格式
                if number_format and isinstance(value, (int, float)):
                    cell.number_format = number_format

        # 自动调整列宽
        self._auto_adjust_column_width(sheet, headers, data)

        # 保存为字节流
        output = io.BytesIO()
        self.workbook.save(output)
        output.seek(0)

        return output.getvalue()

    def export_multi_sheet(
        self,
        sheets_data: list[dict[str, Any]]
    ) -> bytes:
        """
        导出多工作表Excel

        Args:
            sheets_data: 工作表数据列表，每个元素包含:
                - name: 工作表名称
                - data: 数据列表
                - headers: 表头列表

        Returns:
            Excel文件字节流
        """
        self.workbook = Workbook()

        # 删除默认工作表
        default_sheet = self.workbook.active
        self.workbook.remove(default_sheet)

        for sheet_info in sheets_data:
            sheet_name = sheet_info.get("name", "Sheet")
            data = sheet_info.get("data", [])
            headers = sheet_info.get("headers", [])

            sheet = self.workbook.create_sheet(title=sheet_name)

            # 写入表头
            for col_idx, header in enumerate(headers, 1):
                cell = sheet.cell(row=1, column=col_idx, value=header)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

            # 写入数据
            for row_idx, row_data in enumerate(data, 2):
                for col_idx, header in enumerate(headers, 1):
                    value = row_data.get(header) if isinstance(row_data, dict) else None
                    if isinstance(value, datetime):
                        value = value.strftime("%Y-%m-%d %H:%M:%S")
                    elif value is None:
                        value = ""
                    sheet.cell(row=row_idx, column=col_idx, value=value)

            self._auto_adjust_column_width(sheet, headers, data)

        # 保存为字节流
        output = io.BytesIO()
        self.workbook.save(output)
        output.seek(0)

        return output.getvalue()

    def _auto_adjust_column_width(
        self,
        sheet,
        headers: list[str],
        data: list[dict[str, Any]]
    ) -> None:
        """
        自动调整列宽

        Args:
            sheet: 工作表对象
            headers: 表头列表
            data: 数据列表
        """
        for col_idx, header in enumerate(headers, 1):
            max_length = len(header)

            for row_data in data:
                value = row_data.get(header) if isinstance(row_data, dict) else None
                if value:
                    # 计算字符串长度
                    if isinstance(value, datetime):
                        value = value.strftime("%Y-%m-%d %H:%M:%S")
                    max_length = max(max_length, len(str(value)))

            # 设置列宽（添加一些缓冲空间）
            adjusted_width = min(max_length + 2, 50)  # 最大宽度限制为50
            sheet.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    def generate_filename(self, report_type: str) -> str:
        """
        生成Excel文件名

        Args:
            report_type: 报表类型

        Returns:
            文件名
        """
        timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
        return f"{report_type}_report_{timestamp}.xlsx"

    def get_mime_type(self) -> str:
        """
        获取Excel MIME类型

        Returns:
            MIME类型字符串
        """
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_file_extension(self) -> str:
        """
        获取Excel文件扩展名

        Returns:
            文件扩展名
        """
        return ".xlsx"

    def export_report_data(
        self,
        report_type: str,
        data: dict[str, Any]
    ) -> bytes:
        """
        根据报表类型导出数据

        Args:
            report_type: 报表类型 (energy/trend/alarm/runtime)
            data: 报表数据

        Returns:
            Excel文件字节流
        """
        if report_type == "energy":
            return self._export_energy_report(data)
        elif report_type == "trend":
            return self._export_trend_report(data)
        elif report_type == "alarm":
            return self._export_alarm_report(data)
        elif report_type == "runtime":
            return self._export_runtime_report(data)
        else:
            raise ValueError(f"不支持的报表类型: {report_type}")

    def _export_energy_report(self, data: dict[str, Any]) -> bytes:
        """导出能耗报表"""
        headers = ["设备ID", "设备名称", "总能耗(kWh)", "平均功率(W)", "最大功率(W)", "运行时长(小时)"]
        rows = []
        for item in data.get("data", []):
            rows.append({
                "设备ID": item.get("device_id"),
                "设备名称": item.get("device_name"),
                "总能耗(kWh)": item.get("total_energy"),
                "平均功率(W)": item.get("avg_power"),
                "最大功率(W)": item.get("max_power"),
                "运行时长(小时)": item.get("runtime_hours")
            })

        return self.export_data(rows, headers, sheet_name="能耗统计", number_format="#,##0.00")

    def _export_trend_report(self, data: dict[str, Any]) -> bytes:
        """导出温湿度趋势报表"""
        sheets_data = []

        for item in data.get("data", []):
            device_id = item.get("device_id")
            device_name = item.get("device_name")

            # 温度数据
            temp_rows = []
            for point in item.get("temp_trend", []):
                temp_rows.append({
                    "时间": point.get("time"),
                    "温度": point.get("value")
                })

            # 湿度数据
            humi_rows = []
            for point in item.get("humi_trend", []):
                humi_rows.append({
                    "时间": point.get("time"),
                    "湿度": point.get("value")
                })

            sheets_data.append({
                "name": f"{device_name}_温度",
                "headers": ["时间", "温度"],
                "data": temp_rows
            })
            sheets_data.append({
                "name": f"{device_name}_湿度",
                "headers": ["时间", "湿度"],
                "data": humi_rows
            })

        return self.export_multi_sheet(sheets_data)

    def _export_alarm_report(self, data: dict[str, Any]) -> bytes:
        """导出告警报表"""
        headers = ["告警类型", "严重程度", "数量", "已解决", "未解决"]
        rows = []
        for item in data.get("data", []):
            rows.append({
                "告警类型": item.get("type"),
                "严重程度": item.get("severity"),
                "数量": item.get("count"),
                "已解决": item.get("resolved_count"),
                "未解决": item.get("unresolved_count")
            })

        return self.export_data(rows, headers, sheet_name="告警统计")

    def _export_runtime_report(self, data: dict[str, Any]) -> bytes:
        """导出运行时长报表"""
        headers = ["设备ID", "设备名称", "分区", "运行时长(小时)", "开机占比(%)", "开机次数", "关机次数"]
        rows = []
        for item in data.get("data", []):
            rows.append({
                "设备ID": item.get("device_id"),
                "设备名称": item.get("device_name"),
                "分区": item.get("zone_name") or "",
                "运行时长(小时)": item.get("total_runtime_hours"),
                "开机占比(%)": item.get("on_time_percentage"),
                "开机次数": item.get("on_count"),
                "关机次数": item.get("off_count")
            })

        return self.export_data(rows, headers, sheet_name="运行时长", number_format="#,##0.00")